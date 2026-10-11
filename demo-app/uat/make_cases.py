"""Write the Lumen Clinic UAT workbooks in the kit's column layout.

  lumen-clinic-tests.xlsx      first pack
  lumen-clinic-tests-v2.xlsx   a later edit: wording, a real change, a removal, a near-duplicate

Run from anywhere: python demo-app/uat/make_cases.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
START = "http://localhost:4173"
EMPLOYEE = "Employee test account"
MANAGER = "Manager test account"

COLUMNS = [
    ("test_id", "Test ID *", 12),
    ("group", "Group *", 16),
    ("title", "Title *", 36),
    ("start_page", "Start page *", 36),
    ("login_as", "Login as", 24),
    ("preconditions", "Preconditions", 36),
    ("steps", "Steps *", 54),
    ("expected_results", "Expected results *", 54),
    ("automation", "Automation", 14),
    ("priority", "Priority", 12),
    ("notes", "Notes", 42),
]

V1 = [
    {
        "test_id": "TC-01", "group": "Today", "title": "Employee sees the morning board",
        "start_page": START + "/", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace has the 09:00 appointment.",
        "steps": '1. Click "Today"',
        "expected_results": '1. The heading shows "Today"\n1. The board shows "Signed in as Nora Ellis"\n1. The board shows "Ada Lovelace"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-02", "group": "Today", "title": "Manager sees their own name",
        "start_page": START + "/", "login_as": MANAGER, "preconditions": "",
        "steps": '1. Click "Today"',
        "expected_results": '1. The board shows "Signed in as Samir Haddad"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-03", "group": "Billing", "title": "Employee cannot open billing",
        "start_page": START + "/", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "Billing"',
        "expected_results": '1. The page shows "You do not have access to billing."',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-04", "group": "Patients", "title": "Add a patient",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Two patients are already on file.",
        "steps": '1. Click "Add patient"\n2. Type "Mina Chen" in Full name\n3. Type "1008" in MRN\n4. Type "1994-03-12" in Date of birth\n5. Click "Save patient"',
        "expected_results": '5. The status message shows "Added patient Mina Chen"\n5. The patient list shows "Mina Chen"\n5. The patient list shows "3 patients"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-05", "group": "Patients", "title": "Reject a duplicate MRN",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is already MRN 1001.",
        "steps": '1. Click "Add patient"\n2. Type "Copy Ada" in Full name\n3. Type "1001" in MRN\n4. Type "1980-01-01" in Date of birth\n5. Click "Save patient"',
        "expected_results": '5. The status message shows "MRN 1001 is already on file."\n5. The patient list shows "2 patients"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-06", "group": "Patients", "title": "Search finds one patient",
        "start_page": START + "/patients", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Type "Turing" in Search patients',
        "expected_results": '1. The patient list shows "Alan Turing"\n1. The patient list shows "1 patient"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-07", "group": "Patients", "title": "Open Ada Lovelace's record",
        "start_page": START + "/patients", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "Ada Lovelace"',
        "expected_results": '1. The heading shows "Ada Lovelace"\n1. The record shows "MRN 1001"\n1. The record shows "Follow-up: not set"\n1. The record shows "Phone: 555-0148"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-08", "group": "Appointments", "title": "Book an afternoon visit",
        "start_page": START + "/appointments", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is on file.",
        "steps": '1. Click "Book appointment"\n2. Choose "Ada Lovelace" in Patient\n3. Type "2026-10-12" in Date\n4. Type "14:30" in Time\n5. Type "Dr. Idris" in Clinician\n6. Click "Save appointment"',
        "expected_results": '6. The status message shows "Booked Ada Lovelace at 14:30"\n6. The appointment list shows "14:30"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-09", "group": "Appointments", "title": "Delete the morning appointment",
        "start_page": START + "/appointments", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace has an appointment at 09:00.",
        "steps": '1. Click "Delete 09:00 Ada Lovelace"\n2. Click "Delete appointment"',
        "expected_results": '2. The status message shows "Deleted the 09:00 appointment for Ada Lovelace"',
        "automation": "auto", "priority": "High",
        "notes": "The delete buttons ask a person before the click. That is part of this flow.",
    },
    {
        "test_id": "TC-10", "group": "Prescriptions", "title": "Employee cannot prescribe",
        "start_page": START + "/prescriptions", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "New prescription"',
        "expected_results": '1. The status message shows "Only a manager can prescribe."',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-11", "group": "Prescriptions", "title": "Manager prescribes paracetamol",
        "start_page": START + "/prescriptions", "login_as": MANAGER, "preconditions": "",
        "steps": '1. Click "New prescription"\n2. Choose "Alan Turing" in Patient\n3. Type "Paracetamol" in Medicine\n4. Type "500 mg" in Dose\n5. Click "Save prescription"',
        "expected_results": '5. The status message shows "Prescribed Paracetamol for Alan Turing"\n5. The prescription list shows "Paracetamol"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-12", "group": "Prescriptions", "title": "A prescription needs a dose",
        "start_page": START + "/prescriptions", "login_as": MANAGER, "preconditions": "",
        "steps": '1. Click "New prescription"\n2. Choose "Alan Turing" in Patient\n3. Type "Paracetamol" in Medicine\n4. Click "Save prescription"',
        "expected_results": '4. The status message shows "Enter a dose."',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-13", "group": "Billing", "title": "Manager pays invoice INV-100",
        "start_page": START + "/billing", "login_as": MANAGER,
        "preconditions": "Invoice INV-100 is unpaid.",
        "steps": '1. Click "Pay INV-100"\n2. Type "Card" in Method\n3. Click "Pay invoice"',
        "expected_results": '3. The status message shows "Paid INV-100"\n3. The invoice list shows "Paid"',
        "automation": "auto", "priority": "High",
        "notes": "Pay asks a person before the click.",
    },
    {
        "test_id": "TC-14", "group": "Stock", "title": "Paracetamol is low",
        "start_page": START + "/stock", "login_as": EMPLOYEE, "preconditions": "Paracetamol has 4 on hand.",
        "steps": '1. Click "Stock"',
        "expected_results": '1. The stock list shows "Paracetamol"\n1. The stock list shows "Low"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-15", "group": "Stock", "title": "Manager sets paracetamol to 8",
        "start_page": START + "/stock", "login_as": MANAGER, "preconditions": "",
        "steps": '1. Click "Adjust Paracetamol"\n2. Type "8" in On hand\n3. Click "Save stock"',
        "expected_results": '3. The status message shows "Paracetamol is now 8 on hand"\n3. The stock list shows "8"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-16", "group": "Reports", "title": "Unpaid total on the clinic day",
        "start_page": START + "/reports", "login_as": EMPLOYEE, "preconditions": "INV-100 is unpaid and INV-101 is paid.",
        "steps": '1. Click "Reports"',
        "expected_results": '1. The heading shows "Reports"\n1. The report shows "Unpaid €40.00"',
        "automation": "auto", "priority": "Low", "notes": "",
    },
    {
        "test_id": "TC-17", "group": "Today", "title": "Sign out",
        "start_page": START + "/", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "Sign out"',
        "expected_results": '1. The heading shows "Sign in"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-18", "group": "Visits", "title": "Follow-up looks right",
        "start_page": START + "/patients", "login_as": EMPLOYEE, "preconditions": "",
        "steps": "1. Open a patient\n2. Check the follow-up",
        "expected_results": "2. The follow-up is correct",
        "automation": "auto", "priority": "Low",
        "notes": "Deliberately vague. Clarification should ask which patient and what correct means.",
    },
    {
        "test_id": "TC-19", "group": "Mail", "title": "Patient confirms the visit by phone",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace's record shows Phone: 555-0148.",
        "steps": "1. Call the number on Ada Lovelace's record\n2. Ask her to confirm the 09:00 visit",
        "expected_results": "2. Ada Lovelace confirms the 09:00 visit",
        "automation": "manual", "priority": "Medium",
        "notes": "A person makes the call. The agents do not run this case. Record the result at sign-off.",
    },
    {
        "test_id": "TC-21", "group": "Appointments", "title": "Check in the morning visit",
        "start_page": START + "/appointments", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is Booked at 09:00. Nobody is in the waiting room.",
        "steps": '1. Click "Check in 09:00 Ada Lovelace"\n2. Click "Today"',
        "expected_results": '1. The status message shows "Ada Lovelace is checked in"\n1. The appointment list shows "Arrived"\n2. The board shows "Waiting: Ada Lovelace"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-22", "group": "Appointments", "title": "Refuse a second booking at the same time",
        "start_page": START + "/appointments", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is already booked at 09:00 on 2026-10-11.",
        "steps": '1. Click "Book appointment"\n2. Choose "Ada Lovelace" in Patient\n3. Type "2026-10-11" in Date\n4. Type "09:00" in Time\n5. Type "Dr. Okonkwo" in Clinician\n6. Click "Save appointment"',
        "expected_results": '6. The status message shows "Ada Lovelace is already booked at 09:00."',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-23", "group": "Appointments", "title": "Refuse a clinician who is already booked",
        "start_page": START + "/appointments", "login_as": EMPLOYEE,
        "preconditions": "Dr. Idris is already booked at 09:00 on 2026-10-11.",
        "steps": '1. Click "Book appointment"\n2. Choose "Alan Turing" in Patient\n3. Type "2026-10-11" in Date\n4. Type "09:00" in Time\n5. Type "Dr. Idris" in Clinician\n6. Click "Save appointment"',
        "expected_results": '6. The status message shows "Dr. Idris is already booked at 09:00."',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-24", "group": "Patients", "title": "Ada's penicillin allergy is on the record",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is allergic to Penicillin.",
        "steps": '1. Click "Ada Lovelace"',
        "expected_results": '1. The record shows "Allergy: Penicillin"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-25", "group": "Prescriptions", "title": "Penicillin is refused for Ada",
        "start_page": START + "/prescriptions", "login_as": MANAGER,
        "preconditions": "Ada Lovelace is allergic to Penicillin.",
        "steps": '1. Click "New prescription"\n2. Choose "Ada Lovelace" in Patient\n3. Type "Penicillin" in Medicine\n4. Type "250 mg" in Dose\n5. Click "Save prescription"',
        "expected_results": '5. The status message shows "Ada Lovelace is allergic to Penicillin."\n5. The prescription list shows "No prescriptions this session."',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-26", "group": "Patients", "title": "Plan a follow-up",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace's follow-up is not set.",
        "steps": '1. Click "Ada Lovelace"\n2. Click "Plan follow-up"\n3. Type "2026-11-02" in Follow-up date\n4. Click "Save follow-up"',
        "expected_results": '4. The status message shows "Follow-up for Ada Lovelace is 2026-11-02"\n4. The record shows "Follow-up: 2026-11-02"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-27", "group": "Patients", "title": "Save a clinical note",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace has no notes this session.",
        "steps": '1. Click "Ada Lovelace"\n2. Click "Add note"\n3. Type "Sore throat since Monday" in Note\n4. Click "Save note"',
        "expected_results": '4. The status message shows "Saved a note for Ada Lovelace"\n4. The notes show "Sore throat since Monday"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-28", "group": "Patients", "title": "Remove the penicillin allergy",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is allergic to Penicillin.",
        "steps": '1. Click "Ada Lovelace"\n2. Click "Remove Penicillin"\n3. Click "Remove allergy"',
        "expected_results": '3. The status message shows "Removed Penicillin for Ada Lovelace"\n3. The record shows "Allergy: none"',
        "automation": "auto", "priority": "High",
        "notes": "Remove asks a person before the click, and the dialog asks again.",
    },
    {
        "test_id": "TC-29", "group": "Referrals", "title": "Employee cannot send a referral",
        "start_page": START + "/referrals", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "New referral"',
        "expected_results": '1. The status message shows "Only a manager can send a referral."',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-30", "group": "Referrals", "title": "Manager sends Alan to cardiology",
        "start_page": START + "/referrals", "login_as": MANAGER, "preconditions": "",
        "steps": '1. Click "New referral"\n2. Choose "Alan Turing" in Patient\n3. Type "Cardiology" in Specialty\n4. Type "Chest pain" in Reason\n5. Click "Send referral"',
        "expected_results": '5. The status message shows "Sent a referral for Alan Turing to Cardiology"\n5. The referral list shows "Cardiology"',
        "automation": "auto", "priority": "High",
        "notes": "Send asks a person before the click.",
    },
    {
        "test_id": "TC-31", "group": "Labs", "title": "Order a lipid panel",
        "start_page": START + "/labs", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "Order lab"\n2. Choose "Alan Turing" in Patient\n3. Type "Lipid panel" in Test\n4. Click "Save lab"',
        "expected_results": '4. The status message shows "Ordered Lipid panel for Alan Turing"\n4. The lab list shows "Lipid panel"',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-32", "group": "Labs", "title": "Manager records the blood count",
        "start_page": START + "/labs", "login_as": MANAGER,
        "preconditions": "Full blood count for Ada Lovelace is pending.",
        "steps": '1. Click "Record Full blood count"\n2. Type "Normal" in Result\n3. Click "Save result"',
        "expected_results": '3. The status message shows "Recorded Normal for Full blood count"\n3. The lab list shows "Normal"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-33", "group": "Labs", "title": "Employee cannot record a result",
        "start_page": START + "/labs", "login_as": EMPLOYEE,
        "preconditions": "Full blood count for Ada Lovelace is pending.",
        "steps": '1. Click "Record Full blood count"',
        "expected_results": '1. The status message shows "Only a manager can record a result."',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-34", "group": "Patients", "title": "Search with no match",
        "start_page": START + "/patients", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Type "Zzz" in Search patients',
        "expected_results": '1. The patient list shows "No patient matches."\n1. The patient list shows "0 patients"',
        "automation": "auto", "priority": "Low", "notes": "",
    },
    {
        "test_id": "TC-35", "group": "Billing", "title": "Raise an invoice and see it on the report",
        "start_page": START + "/billing", "login_as": MANAGER,
        "preconditions": "INV-100 is unpaid for €40.00. INV-102 does not exist yet.",
        "steps": '1. Click "New invoice"\n2. Choose "Ada Lovelace" in Patient\n3. Type "25.00" in Amount\n4. Click "Save invoice"\n5. Click "Reports"',
        "expected_results": '4. The status message shows "Raised INV-102 for €25.00"\n4. The invoice list shows "INV-102"\n5. The report shows "Unpaid €65.00"\n5. The report shows "2 unpaid invoices"',
        "automation": "auto", "priority": "High", "notes": "",
    },
    {
        "test_id": "TC-36", "group": "Patients", "title": "A patient needs a name",
        "start_page": START + "/patients", "login_as": EMPLOYEE, "preconditions": "",
        "steps": '1. Click "Add patient"\n2. Click "Save patient"',
        "expected_results": '2. The status message shows "Enter a full name."',
        "automation": "auto", "priority": "Medium", "notes": "",
    },
    {
        "test_id": "TC-37", "group": "Waiting", "title": "The waiting room looks right",
        "start_page": START + "/appointments", "login_as": EMPLOYEE, "preconditions": "",
        "steps": "1. Look at who is waiting\n2. Decide if the room is right",
        "expected_results": "2. The waiting room is correct",
        "automation": "auto", "priority": "Low",
        "notes": "Deliberately vague. Clarification should ask who is waiting and what correct means.",
    },
    {
        "test_id": "TC-38", "group": "Pharmacy", "title": "Confirm penicillin with the pharmacy",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is allergic to Penicillin.",
        "steps": "1. Call the pharmacy about Ada Lovelace\n2. Ask them to confirm she must not be given penicillin",
        "expected_results": "2. The pharmacy confirms Ada Lovelace must not receive penicillin",
        "automation": "manual", "priority": "Medium",
        "notes": "A person makes the call. The agents do not run this case. Record the result at sign-off.",
    },
]


def write(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    keys = [k for k, _, _ in COLUMNS]
    ws.append([header for _, header, _ in COLUMNS])
    for col, (_, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1A3A34", end_color="1A3A34")
        ws.column_dimensions[cell.column_letter].width = width
    for row in rows:
        ws.append([row.get(k, "") for k in keys])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for key, values in (("priority", "High,Medium,Low"), ("automation", "auto,manual")):
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = ws.cell(1, keys.index(key) + 1).column_letter
        dv.add(f"{letter}2:{letter}200")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def v2() -> list[dict]:
    rows = [dict(r) for r in V1 if r["test_id"] != "TC-14"]
    for row in rows:
        if row["test_id"] == "TC-04":
            row["steps"] = row["steps"].replace('Click "Save patient"', 'click "Save patient"')
        if row["test_id"] == "TC-08":
            row["steps"] = row["steps"].replace("14:30", "15:00")
            row["expected_results"] = row["expected_results"].replace("14:30", "15:00")
    rows.append({
        "test_id": "TC-20", "group": "Patients", "title": "An MRN already on file is refused",
        "start_page": START + "/patients", "login_as": EMPLOYEE,
        "preconditions": "Ada Lovelace is already MRN 1001.",
        "steps": '1. Click "Add patient"\n2. Type "Copy Ada" in Full name\n3. Type "1001" in MRN\n4. Type "1980-01-01" in Date of birth\n5. Click "Save patient"',
        "expected_results": '5. The status message shows "MRN 1001 is already on file."\n5. The patient list shows "2 patients"',
        "automation": "auto", "priority": "Medium",
        "notes": "New case, close to TC-05, so change detection can call it similar.",
    })
    return rows


if __name__ == "__main__":
    write(ROOT / "lumen-clinic-tests.xlsx", V1)
    write(ROOT / "lumen-clinic-tests-v2.xlsx", v2())
    print(f"Wrote {len(V1)} cases and {len(v2())} changed cases in {ROOT}")
