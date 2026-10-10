import asyncio
import json


def test_health(client, kit_dir):
    r = client.get("/api/health").json()
    assert r["ok"] and r["kit_found"] and r["kit"] == str(kit_dir)


def test_api_key_is_never_returned(client):
    r = client.put("/api/settings", json={"api_key": "sk-ant-test-1234567890abcd"}).json()
    assert r["api_key_set"] and r["api_key_hint"] == "…abcd"
    assert "sk-ant" not in json.dumps(client.get("/api/settings").json())
    r = client.put("/api/settings", json={"clear_api_key": True, "models": {"uat-clarifier": "claude-opus-5-5"}}).json()
    assert not r["api_key_set"] and r["models"]["uat-clarifier"] == "claude-opus-5-5"


def test_feature_switches_cascade_and_restore(client):
    data = client.get("/api/features").json()
    assert not data["problems"] and len(data["items"]) > 60
    preview = client.post("/api/features/switch", json={"id": "step-screenshots", "action": "off"}).json()
    assert preview["allowed"] and "skill-comparing-screens" in preview["also"] and not preview["applied"]
    client.post("/api/features/switch", json={"id": "step-screenshots", "action": "off", "apply": True})
    items = {i["id"]: i for i in client.get("/api/features").json()["items"]}
    assert not items["step-screenshots"]["on"] and not items["agent-screen-reviewer"]["on"]
    client.post("/api/features/switch", json={"id": "agent-screen-reviewer", "action": "on", "apply": True})
    items = {i["id"]: i for i in client.get("/api/features").json()["items"]}
    assert items["step-screenshots"]["on"] and items["agent-screen-reviewer"]["on"]
    core = client.post("/api/features/switch", json={"id": "stage-run", "action": "off"}).json()
    assert core["allowed"] is False


def test_env_values_are_write_only(client, kit_dir):
    key = client.post("/api/env/accounts", json={"label": "Manager test account"}).json()["key"]
    assert key == "MANAGER_TEST_ACCOUNT"
    client.put("/api/env", json={"key": "MANAGER_TEST_ACCOUNT_PASSWORD", "value": "s3cret-value"})
    body = json.dumps(client.get("/api/env").json())
    assert "s3cret-value" not in body
    acc = {a["name"]: a for a in client.get("/api/env").json()["accounts"]}
    assert acc["MANAGER_TEST_ACCOUNT"]["password_set"] and not acc["MANAGER_TEST_ACCOUNT"]["user_set"]
    assert client.put("/api/env", json={"key": "PATH", "value": "x"}).status_code == 400


def test_guard_refuses_env_and_other_sites(kit_dir):
    from app import guard
    assert guard._decide_blocking("search/textSearch", {"query": ".env"})[0] == "deny"
    assert guard._decide_blocking("playwright-test/browser_navigate", {"url": "https://evil.example.com"})[0] == "deny"
    assert guard._decide_blocking("playwright-test/planner_save_plan", {})[0] == "deny"
    assert guard._decide_blocking("playwright-test/browser_click", {"element": "Delete Taxi"})[0] == "ask"


def test_new_run_intake_and_plan_without_ai(client, kit_dir):
    with open(kit_dir / "examples" / "expense-tracker-tests.xlsx", "rb") as f:
        r = client.post("/api/runs", files={"file": ("expense-tracker-tests.xlsx", f)},
                        data={"strictness": "standard", "start": "false"}).json()
    run_id = r["id"]
    from app.engine import Ctx, engine
    c = Ctx(run_id)
    c.stage = "intake"
    result, note = asyncio.run(engine.stage_intake(c))
    assert result == "done" and (kit_dir / "runs" / run_id / "cases.json").exists()
    assert len(engine.cases_to_clarify(c)) >= 5
    c.stage = "plan"
    result, note = asyncio.run(engine.stage_plan(c))
    assert result == "done" and "To generate" in note
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["data"]["generate_list"] and [s["id"] for s in run["stages"]][0] == "preflight"
    assert any(x["id"] == run_id for x in client.get("/api/runs").json())


def test_clarify_without_api_key_waits_for_the_person(client, kit_dir):
    run_id = client.get("/api/runs").json()[0]["id"]
    from app.engine import Ctx, engine
    from app.llm import MissingApiKey
    c = Ctx(run_id)
    c.stage = "clarify"
    try:
        asyncio.run(engine.stage_clarify(c))
        raised = False
    except MissingApiKey:
        raised = True
    assert raised


def test_knowledge_endpoints(client):
    assert "Glossary" in client.get("/api/kit/glossary").json()["text"]
    assert any(c["name"] == "guardrails.yaml" for c in client.get("/api/kit/config").json())
    assert client.put("/api/kit/config/guardrails.yaml", json={"text": "a: [unclosed"}).status_code == 400
    assert client.put("/api/kit/config/profile.yaml", json={"text": "switched_off: []"}).status_code == 400
    assert any(a["name"] == "playwright-test-planner" for a in client.get("/api/kit/agents").json())
    assert len(client.get("/api/kit/skills").json()) >= 17
    assert client.get("/api/kit/text", params={"path": ".env"}).status_code == 400
    assert client.get("/api/kit/template").status_code == 200


def _latest_run(client):
    return client.get("/api/runs").json()[0]["id"]


def test_visual_regression_groups_screens_per_test_with_ai_comments(client, kit_dir):
    run_id = _latest_run(client)
    run = kit_dir / "runs" / run_id
    base = "runs/2026-01-01-0000/evidence/screenshots"
    (run / "screens.json").write_text(json.dumps({
        "baseline_run": "2026-01-01-0000", "compared": 2, "ai_review": "needed", "flagged": [], "pairs": [
            {"name": "TC-01-step3.png", "test_id": "TC-01", "step": 3, "status": "changed", "old": f"{base}/TC-01-step3.png",
             "new": f"runs/{run_id}/evidence/screenshots/TC-01-step3.png", "diff": "x-diff.png", "changed_pixels": 1022},
            {"name": "TC-01-step1.png", "test_id": "TC-01", "step": 1, "status": "unchanged", "old": "a", "new": "b",
             "changed_pixels": 0},
            {"name": "TC-02-step1.png", "test_id": "TC-02", "step": 1, "status": "new", "old": "", "new": "c"}]}),
        encoding="utf-8")
    (run / "screens-review.json").write_text(json.dumps({"TC-01-step3.png": {"label": "unexplained", "note": "Total moved."}}),
                                             encoding="utf-8")
    v = client.get(f"/api/runs/{run_id}/visual").json()
    assert v["baseline_run"] == "2026-01-01-0000"
    first = v["tests"][0]
    assert first["test_id"] == "TC-01" and first["counts"] == {"changed": 1, "unchanged": 1, "new": 0}
    assert [s["step"] for s in first["steps"]] == [1, 3]                  # in step order
    step3 = first["steps"][1]
    assert step3["ai"] == {"label": "unexplained", "note": "Total moved."} and step3["step_text"] == 'Click "Add expense"'
    assert step3["expected"] and v["tests"][1]["counts"]["new"] == 1


def test_clarified_excel_round_trips_through_the_kit_reader(client, kit_dir):
    run_id = _latest_run(client)
    r = client.get(f"/api/runs/{run_id}/clarified.xlsx")
    assert r.status_code == 200
    out = kit_dir / "runs" / run_id / "clarified-cases.xlsx"
    from openpyxl import load_workbook
    ws = load_workbook(out).worksheets[0]
    headers = [c.value for c in ws[1]]
    assert headers[0].startswith("Test ID") and ws.max_row >= 6
    steps_col = next(i for i, h in enumerate(headers) if h.startswith("Steps"))
    assert ws.cell(2, steps_col + 1).value.startswith("1. ")


def test_known_noise_is_added_by_a_person(client, kit_dir):
    assert client.post("/api/kit/console-noise", json={"kind": "console", "pattern": "x", "reason": ""}).status_code == 400
    r = client.post("/api/kit/console-noise", json={"kind": "network", "pattern": "/analytics/", "reason": "blocked in UAT",
                                                     "status": [404]}).json()
    assert r["ok"]
    import yaml
    text = (kit_dir / "config" / "console-noise.yaml").read_text(encoding="utf-8")
    assert text.startswith("# Known harmless")                            # the header comment is kept
    assert {"url_pattern": "/analytics/", "status": [404], "reason": "blocked in UAT"} in yaml.safe_load(text)["network"]


def test_open_trace_refuses_anything_but_a_trace_of_the_run(client):
    run_id = _latest_run(client)
    assert client.post(f"/api/runs/{run_id}/open-trace", json={"path": "../../.env"}).status_code == 400
    assert client.post(f"/api/runs/{run_id}/open-trace", json={"path": "cases.json"}).status_code == 400
