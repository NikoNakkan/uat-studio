"""Read UAT test cases (one row per test case) from Excel into cases.json.

Usage: python read_excel.py <input.xlsx> <output cases.json>
Also writes cases-original.json next to it: what the Excel said, never edited. Clarification
edits cases.json; change detection compares the originals.
Exit 0 = read, 1 = the layout is wrong (printed), 2 = usage or file error.
Content rules (numbering, required values) are checked by validate.py.
"""
import json
import re
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
TEMPLATE = yaml.safe_load((ROOT / "config" / "excel-template.yaml").read_text(encoding="utf-8"))
NUMBERED = re.compile(r"^\s*(\d+)\s*[.)]\s*(.+?)\s*$")


def norm(text):
    return re.sub(r"[\s_\-*]+", "", str(text or "")).lower()


def lines(value):
    return [l for l in str(value or "").replace("\r", "").split("\n") if l.strip()]


def numbered(value, key):
    out = []
    for line in lines(value):
        m = NUMBERED.match(line)
        out.append({key: int(m.group(1)), "text": m.group(2)} if m else {key: None, "text": line.strip()})
    return out


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    if not src.exists():
        print(f"Input file not found: {src}")
        return 2

    wb = load_workbook(src, read_only=True, data_only=True)
    sheet = wb[TEMPLATE["sheet"]] if TEMPLATE["sheet"] in wb.sheetnames else wb.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))

    columns = TEMPLATE["columns"]
    wanted = {norm(c["header"]): key for key, c in columns.items()}
    header_row = next((i for i, r in enumerate(rows) if norm(columns["test_id"]["header"]) in map(norm, r)), None)
    if header_row is None:
        print(f"No header row with '{columns['test_id']['header']}' in sheet '{sheet.title}'.")
        return 1
    position = {wanted[norm(h)]: i for i, h in enumerate(rows[header_row]) if norm(h) in wanted}
    missing = [c["header"] for k, c in columns.items() if c["required"] and k not in position]
    if missing:
        print(f"Missing required column(s): {', '.join(missing)}")
        return 1

    def cell(row, key):
        i = position.get(key)
        value = row[i] if i is not None and i < len(row) else None
        return "" if value is None else str(value).strip()

    cases = []
    for excel_row, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        if not any(v not in (None, "") for v in row):
            continue
        case = {key: cell(row, key) for key in columns if key not in ("steps", "expected_results")}
        case["steps"] = numbered(cell(row, "steps"), "n")
        case["expected_results"] = numbered(cell(row, "expected_results"), "step")
        case["row"] = excel_row
        cases.append(case)

    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps(cases, indent=2, ensure_ascii=False), encoding="utf-8")
    dst.with_name("cases-original.json").write_text(json.dumps(cases, indent=2, ensure_ascii=False), encoding="utf-8")
    groups = sorted({c["group"] for c in cases if c["group"]})
    print(f"Read {len(cases)} test case(s) in {len(groups)} group(s) from sheet '{sheet.title}' -> {dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
