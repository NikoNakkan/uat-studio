"""Build templates/uat-test-cases.xlsx: the agreed template (one row per test case) with an example row
and a "How to fill in" sheet. Column names come from config/excel-template.yaml.

Usage: python templates/make_template.py
"""
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
COLUMNS = yaml.safe_load((ROOT / "config" / "excel-template.yaml").read_text(encoding="utf-8"))["columns"]
WIDTHS = {"test_id": 10, "group": 14, "title": 30, "start_page": 30, "login_as": 22, "preconditions": 24,
          "steps": 46, "expected_results": 46, "automation": 12, "priority": 10, "notes": 30}
HELP = {
    "test_id": "Unique, never reused, e.g. TC-01. Letters, digits and - only.",
    "group": "Cases signed off together, e.g. Expenses. Becomes the test folder.",
    "title": "Short name of what is tested.",
    "start_page": "Full address where the test starts, e.g. https://test.app.com/expenses",
    "login_as": "Name of a test account, e.g. Employee test account. NEVER a password.",
    "preconditions": "What must be true before step 1. Leave empty if nothing.",
    "steps": "One numbered step per line: 1. ...  2. ...  (Alt+Enter for a new line)",
    "expected_results": "One per line, starting with the number of its step: 3. Status shows \"Added Taxi\"",
    "automation": "auto (or empty): the agents run it. manual: a person performs it (email, phone, another system) and records the result at sign-off.",
    "priority": "High, Medium or Low.",
    "notes": "Anything else. Not used by the test.",
}
GUIDE = [
    ("How to fill in the UAT test cases", ""),
    ("", ""),
    ("One row", "Each row is one test case. Do not merge cells or split a case over several rows."),
    ("Steps", "One step per line in the Steps cell, numbered 1., 2., 3. ... in order. Press Alt+Enter for a new line."),
    ("One action per step", "Type, click, choose, open. Name the button or field exactly as it appears on screen, in quotes."),
    ("Values", "Write the exact value to type: Type \"Taxi\" in Description. Not: enter a description."),
    ("Expected results", "One per line. Start each with the number of the step it belongs to: 3. Total shows \"€12.50\"."),
    ("Checkable words", "Say exactly what appears: a text in quotes, an amount, a page. Avoid: works, correct, valid, looks fine."),
    ("Login as", "Name the test account (for example: Employee test account). Never write a password anywhere in this file."),
    ("Automation", "Leave empty (or write auto) for cases the agents run in the browser. Write manual for cases a person must perform: "
                   "an email or SMS arrives, a phone call, a printed document, another system. Manual cases are not generated or run; "
                   "they stay in the results and the person records PASS, FAIL or NOT RUN for them at sign-off."),
    ("Test ID", "Keep a case's Test ID forever, also when you edit it. A new case gets a new ID."),
    ("Group", "Cases in the same group are signed off together."),
    ("Required", "Columns marked * must be filled in."),
]
EXAMPLE = {
    "test_id": "TC-01", "group": "Expenses", "title": "Add an expense",
    "start_page": "http://localhost:4173", "login_as": "", "preconditions": "",
    "steps": '1. Type "Taxi" in Description\n2. Type 12.50 in Amount (EUR)\n3. Click "Add expense"',
    "expected_results": '3. The status message shows "Added Taxi"\n3. The total shows "€12.50"',
    "automation": "auto", "priority": "High", "notes": "Example row: replace it with your own cases.",
}


def build(path, rows, guide=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    keys = list(COLUMNS)
    ws.append([COLUMNS[k]["header"] + (" *" if COLUMNS[k]["required"] else "") for k in keys])
    for i, k in enumerate(keys, start=1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="305496", end_color="305496")
        cell.comment = Comment(HELP[k], "UAT template")
        ws.column_dimensions[cell.column_letter].width = WIDTHS[k]
    for row in rows:
        ws.append([row.get(k, "") for k in keys])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    for key in ("priority", "automation"):                 # drop-down lists from the template's allowed values
        values = COLUMNS[key].get("values")
        if not values:
            continue
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        col = ws.cell(2, keys.index(key) + 1).column_letter
        dv.add(f"{col}2:{col}500")
    if guide:
        g = wb.create_sheet("How to fill in")
        for a, b in GUIDE:
            g.append([a, b])
        g["A1"].font = Font(bold=True, size=14)
        for r in range(3, g.max_row + 1):
            g.cell(r, 1).font = Font(bold=True)
            g.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        g.column_dimensions["A"].width = 22
        g.column_dimensions["B"].width = 100
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"Wrote {path.relative_to(ROOT).as_posix()} ({len(rows)} case(s))")


if __name__ == "__main__":
    build(ROOT / "templates" / "uat-test-cases.xlsx", [EXAMPLE])
