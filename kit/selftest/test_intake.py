"""reading-uat-excel: read_excel.py and validate.py, and the template built by templates/make_template.py."""
from conftest import case

READ = ".github/skills/reading-uat-excel/scripts/read_excel.py"
VALIDATE = ".github/skills/reading-uat-excel/scripts/validate.py"


def test_examples_read_and_validate(kit):
    run = kit.intake("examples/expense-tracker-tests.xlsx")
    cases = kit.json(f"{run}/cases.json")
    assert [c["test_id"] for c in cases] == [f"TC-0{i}" for i in range(1, 7)]
    assert cases[0]["steps"][0] == {"n": 1, "text": 'Type "Taxi" in Description'}
    assert cases[0]["expected_results"][0]["step"] == 3
    assert kit.path(f"{run}/cases-original.json").exists()


def test_template_file_is_readable(kit):
    kit.run("templates/make_template.py", expect=0)
    run = kit.intake("templates/uat-test-cases.xlsx")
    cases = kit.json(f"{run}/cases.json")
    assert len(cases) == 1 and cases[0]["automation"] == "auto"


def test_missing_required_column_is_refused(kit):
    kit.excel("runs/bad.xlsx", [])
    p = kit.path("runs/bad.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb.active
    ws.delete_cols(next(i for i, c in enumerate(ws[1], start=1) if str(c.value).startswith("Steps")))
    wb.save(p)
    run = kit.new_run("runs/bad.xlsx")
    r = kit.run(READ, "runs/bad.xlsx", f"{run}/cases.json", expect=1)
    assert "Missing required column(s): Steps" in r.out


def test_validate_names_row_and_column(kit):
    rows = [case("TC-01", steps="1. Click \"Add\"\n3. Click \"Save\"", expected="2. Something"),
            case("TC-01", title="Duplicate id", start_page="https://elsewhere.example.com/x"),
            case("TC 03", login_as="Nobody account", automation="sometimes", priority="Urgent"),
            dict(case("TC-04"), title="")]
    kit.excel("runs/cases.xlsx", rows)
    run = kit.intake("runs/cases.xlsx", expect_valid=False)
    r = kit.run(VALIDATE, f"{run}/cases.json", expect=1)
    for text in ["Row 2, column 'Steps' (TC-01): steps must be numbered 1, 2, 3",
                 "column 'Expected results' (TC-01): refers to step 2, which does not exist",
                 "the same Test ID is used on more than one row",
                 "https://elsewhere.example.com is not in config/guardrails.yaml allowed_origins",
                 "Row 4, column 'Test ID' (TC 03): use only letters",
                 "'sometimes' is not one of auto, manual",
                 "'Urgent' is not one of High, Medium, Low",
                 "Row 5, column 'Title' (TC-04): required, but empty"]:
        assert text in r.out, text


def test_validate_counts_manual_cases(kit):
    kit.excel("runs/cases.xlsx", [case("TC-01"), case("TC-02", title="Email arrives", automation="manual"),
                                  case("TC-03", title="SMS arrives", automation="Manual")])
    run = kit.intake("runs/cases.xlsx")
    r = kit.run(VALIDATE, f"{run}/cases.json", expect=0)
    assert "2 manual (performed by a person, not generated)" in r.out


def test_excel_without_automation_column_still_reads(kit):
    """Older files have no Automation column: every case counts as auto."""
    kit.excel("runs/old.xlsx", [case("TC-01")])
    p = kit.path("runs/old.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb.active
    ws.delete_cols(next(i for i, c in enumerate(ws[1], start=1) if c.value == "Automation"))
    wb.save(p)
    run = kit.intake("runs/old.xlsx")
    assert kit.json(f"{run}/cases.json")[0]["automation"] == ""
