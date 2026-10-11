"""verifying-evidence, diagnosing-failures and writing-results-excel: from a finished run to a sealed results.xlsx."""
import json

import pytest
from openpyxl import load_workbook

from conftest import GOOD_TEST, case

VERIFY = ".github/skills/verifying-evidence/scripts/verify_evidence.py"
DIAGNOSE = ".github/skills/diagnosing-failures/scripts/record_diagnosis.py"
DECIDE = ".github/skills/writing-results-excel/scripts/record_decision.py"
RESULTS = ".github/skills/writing-results-excel/scripts/write_results.py"


def results_json(kit, run, entries):
    """A Playwright results.json with one trace attachment per test (the trace file is created on disk)."""
    specs = []
    for tid, status in entries:
        trace = kit.write(f"{run}/evidence/traces/{tid}-trace.zip", "zip")
        specs.append({"title": f"{tid} title", "tests": [{"results": [
            {"status": status, "attachments": [{"name": "trace", "path": str(trace)}]}]}]})
    kit.write_json(f"{run}/evidence/results.json", {"suites": [{"specs": specs}]})


@pytest.fixture
def finished_run(kit):
    """TC-01 passed with full evidence, TC-02 failed, TC-03 manual (Expenses); TC-04 manual alone in Validation."""
    kit.excel("runs/cases.xlsx", [
        case("TC-01"), case("TC-02", title="Total of two"),
        case("TC-03", title="Receipt email arrives", automation="manual"),
        case("TC-04", group="Validation", title="Printed report is correct", automation="manual")])
    run = kit.intake("runs/cases.xlsx")
    kit.write("tests/expenses/TC-01-add-an-expense.spec.ts", GOOD_TEST)
    progress = kit.json(f"{run}/progress.json")
    progress["tests"] = {
        "TC-01": {"result": "passed", "file": "tests/expenses/TC-01-add-an-expense.spec.ts", "error": "", "evidence": "evidence"},
        "TC-02": {"result": "failed", "file": "tests/expenses/TC-02-total-of-two.spec.ts", "error": "boom", "evidence": "evidence"},
    }
    kit.write_json(f"{run}/progress.json", progress)
    for n in (1, 2, 3):
        kit.write(f"{run}/evidence/screenshots/TC-01-step{n}.png", "png")
    results_json(kit, run, [("TC-01", "passed"), ("TC-02", "failed")])
    kit.write_json(f"{run}/test-check.json", {"files": {"tests/expenses/TC-01-add-an-expense.spec.ts":
                                                        {"test_id": "TC-01", "ok": True, "problems": []}}})
    return run


def test_recommendations(kit, finished_run):
    run = finished_run
    r = kit.run(VERIFY, run, expect=0)
    assert "Per test: PASS 1, CONCERNS 1, FAIL 0, MANUAL 2" in r.out
    check = kit.json(f"{run}/evidence-check.json")
    assert check["tests"]["TC-01"]["recommendation"] == "PASS" and check["tests"]["TC-01"]["screenshots"] == "3/3"
    assert check["tests"]["TC-02"]["reasons"] == ["failure not diagnosed"]
    assert check["tests"]["TC-03"]["recommendation"] == "MANUAL" and check["tests"]["TC-03"]["manual"] is True
    assert check["units"]["Expenses"]["recommendation"] == "CONCERNS"        # worst of the automated ones
    assert check["units"]["Expenses"]["manual"] == ["TC-03"]
    assert check["units"]["Validation"]["recommendation"] == "MANUAL"        # nothing automated to judge


def test_a_pass_needs_every_piece_of_evidence(kit, finished_run):
    run = finished_run
    kit.path(f"{run}/evidence/screenshots/TC-01-step2.png").unlink()
    kit.run(VERIFY, run, expect=0)
    tc1 = kit.json(f"{run}/evidence-check.json")["tests"]["TC-01"]
    assert tc1["recommendation"] == "CONCERNS" and tc1["missing"] == ["screenshot of step 2"]

    kit.write(f"{run}/evidence/screenshots/TC-01-step2.png", "png")
    kit.write_json(f"{run}/test-check.json", {"files": {}})
    kit.run(VERIFY, run, expect=0)
    assert "no test check recorded" in kit.json(f"{run}/evidence-check.json")["tests"]["TC-01"]["missing"][0]

    kit.run("scripts/features.py", "off", "step-screenshots", "--yes", expect=0)   # switched off: not required
    kit.path(f"{run}/evidence/screenshots/TC-01-step2.png").unlink()
    kit.write_json(f"{run}/test-check.json", {"files": {"x": {"test_id": "TC-01", "ok": True, "problems": []}}})
    r = kit.run(VERIFY, run, expect=0)
    assert kit.json(f"{run}/evidence-check.json")["tests"]["TC-01"]["recommendation"] == "PASS"
    assert "Not required for PASS (switched off" in r.out


def test_failures_are_recommended_from_their_diagnosis(kit, finished_run):
    run = finished_run
    r = kit.run(DIAGNOSE, run, "TC-02", "--cause", "app_broken", expect=0)
    kit.run(VERIFY, run, expect=0)
    tc2 = kit.json(f"{run}/evidence-check.json")["tests"]["TC-02"]
    assert tc2["recommendation"] == "CONCERNS" and "not reproduced yet" in tc2["reasons"][0]
    kit.run(DIAGNOSE, run, "TC-02", "--reproduction", "reproduced", expect=0)
    kit.run(VERIFY, run, expect=0)
    check = kit.json(f"{run}/evidence-check.json")
    assert check["tests"]["TC-02"]["recommendation"] == "FAIL" and check["units"]["Expenses"]["recommendation"] == "FAIL"
    r = kit.run(DIAGNOSE, run, "TC-02", "--repair", "approved", expect=1)
    assert "Only a test labelled test_broken can be repaired" in r.out
    r = kit.run(DIAGNOSE, run, "TC-09", "--cause", "environment", expect=1)
    assert "no result in progress.json" in r.out


def test_sign_off_from_decisions_to_sealed_results(kit, finished_run):
    run = finished_run
    kit.run(VERIFY, run, expect=0)

    r = kit.run(DECIDE, run, "--unit", "Nowhere", "--decision", "PASS", "--by", "Ann", expect=1)
    assert "Unknown unit" in r.out
    r = kit.run(DECIDE, run, "--unit", "Expenses", "--decision", "WAIVED", "--by", "Ann", expect=1)
    assert "needs --reason" in r.out
    r = kit.run(DECIDE, run, "--unit", "Expenses", "--decision", "WAIVED", "--by", "Ann", "--reason", "Known, fix planned", expect=0)
    assert "(the AI recommended CONCERNS)" in r.out
    r = kit.run(DECIDE, run, "--complete", expect=1)
    assert "Not every unit has a decision yet: Validation" in r.out
    kit.run(DECIDE, run, "--unit", "Validation", "--decision", "PASS", "--by", "Ann", expect=0)

    # manual results: required, only for manual cases, never guessed
    r = kit.run(DECIDE, run, "--complete", expect=1)
    assert "Manual case(s) without a result yet: TC-03, TC-04" in r.out
    r = kit.run(DECIDE, run, "--manual", "TC-01", "--result", "PASS", "--by", "Bob", expect=1)
    assert "not a manual case" in r.out
    r = kit.run(DECIDE, run, "--manual", "TC-03", "--result", "PASS", expect=2)
    assert "--manual needs --result and --by" in r.out
    r = kit.run(DECIDE, run, "--manual", "TC-03", "--result", "PASS", "--by", "Bob", "--note", "Email in 2 min",
                "--evidence", "shared/TC-03.png", expect=0)
    assert "1 of 2 manual case(s) recorded" in r.out
    kit.run(DECIDE, run, "--manual", "TC-04", "--result", "NOT_RUN", "--by", "Bob", expect=0)
    manual = kit.json(f"{run}/signoff.json")["manual"]
    assert manual["TC-03"]["evidence"] == "shared/TC-03.png" and manual["TC-04"]["result"] == "NOT_RUN"

    # results before the seal
    r = kit.run(RESULTS, run, expect=0)
    assert "MANUAL 2 (2 with a recorded result); 2 of 2 unit(s) decided" in r.out

    # close and seal
    r = kit.run(DECIDE, run, "--complete", expect=0)
    assert "sealed in manifest.json" in r.out
    manifest = kit.json(f"{run}/manifest.json")["files"]
    assert {"cases.json", "signoff.json", "evidence-check.json", "evidence/screenshots/TC-01-step1.png",
            "tests/expenses/TC-01-add-an-expense.spec.ts"} <= set(manifest)
    r = kit.run(DECIDE, run, "--unit", "Expenses", "--decision", "PASS", "--by", "Ann", expect=1)
    assert "already complete and sealed" in r.out
    r = kit.run(DECIDE, run, "--manual", "TC-03", "--result", "FAIL", "--by", "Bob", expect=1)

    kit.run(RESULTS, run, expect=0)
    wb = load_workbook(kit.path(f"{run}/results.xlsx"))
    assert wb.sheetnames[:3] == ["Run", "Summary", "Tests"]
    run_sheet = {r[0]: r[1] for r in wb["Run"].iter_rows(values_only=True)}
    assert run_sheet["Test cases in the Excel"] == "4 (2 manual, performed by a person)"
    assert run_sheet["Excel SHA-256"] == kit.json(f"{run}/progress.json")["excel_sha256"]
    assert run_sheet["Sites tested"] == "http://localhost:4173" and run_sheet["Sign-off"].startswith("complete, sealed")
    summary = {r[0]: r for r in wb["Summary"].iter_rows(values_only=True)}
    assert summary["Expenses"][1:6] == (3, 1, 1, 0, "PASS 1") and summary["Expenses"][6:8] == ("CONCERNS", "WAIVED")
    assert summary["Validation"][5:8] == ("NOT_RUN 1", "MANUAL", "PASS")
    tests = {r[0]: r for r in wb["Tests"].iter_rows(values_only=True)}
    assert tests["TC-03"][4] == "manual: PASS" and "by Bob" in tests["TC-03"][6] and tests["TC-03"][12] == "shared/TC-03.png"
    assert tests["TC-04"][4] == "manual: NOT_RUN"
    assert tests["TC-01"][4] == "passed" and tests["TC-01"][12] == f"{run}/evidence"
    waivers = list(wb["Waivers"].iter_rows(values_only=True))
    assert waivers[1][:3] == ("Expenses", "Ann", "Known, fix planned")

    # a sealed run whose files changed gives no results
    kit.write("tests/expenses/TC-01-add-an-expense.spec.ts", GOOD_TEST + "\n// edited after the seal\n")
    r = kit.run(RESULTS, run, expect=1)
    assert "Refused: the sign-off was sealed" in r.out and "tests/expenses/TC-01-add-an-expense.spec.ts" in r.out


def test_results_need_the_evidence_check(kit, finished_run):
    r = kit.run(RESULTS, finished_run, expect=2)
    assert "Missing evidence-check.json" in r.out
    r = kit.run(DECIDE, finished_run, "--unit", "Expenses", "--decision", "PASS", "--by", "Ann", expect=1)
    assert "Run verify_evidence.py first" in r.out


def test_sign_off_per_test_unit(kit, finished_run):
    run = finished_run
    kit.write("config/signoff.yaml", kit.read("config/signoff.yaml").replace("sign_off_unit: group", "sign_off_unit: test"))
    kit.run(VERIFY, run, expect=0)
    units = kit.json(f"{run}/evidence-check.json")["units"]
    assert set(units) == {"TC-01", "TC-02", "TC-03", "TC-04"} and units["TC-03"]["recommendation"] == "MANUAL"


def seal(kit, run):
    kit.run(VERIFY, run, expect=0)
    for unit in kit.json(f"{run}/evidence-check.json")["units"]:
        kit.run(DECIDE, run, "--unit", unit, "--decision", "PASS", "--by", "Ann", "--reason", "seen", expect=0)
    for tid in ("TC-03", "TC-04"):
        kit.run(DECIDE, run, "--manual", tid, "--result", "PASS", "--by", "Bob", expect=0)
    kit.run(DECIDE, run, "--complete", expect=0)


def test_stage_bookkeeping_after_the_seal_does_not_break_it(kit, finished_run):
    """Writing results.xlsx is itself a stage: marking it in progress.json must not count as a change."""
    run = finished_run
    seal(kit, run)
    assert kit.json(f"{run}/manifest.json")["files"]["progress.json"].startswith("progress-without-stages:")
    kit.run("scripts/progress.py", "stage", run, "signoff", "started", expect=0)
    kit.run(RESULTS, run, expect=0)


def test_results_embed_the_screenshots_and_explain_empty_sheets(kit, finished_run):
    run = finished_run
    seal(kit, run)
    kit.run(RESULTS, run, expect=0)
    wb = load_workbook(kit.path(f"{run}/results.xlsx"))
    shots = wb["Screenshots"]
    rows = list(shots.iter_rows(values_only=True))[1:]
    assert rows and all("screenshots/" in str(r[6]) for r in rows)  # every step row names its screenshot
    # the self-test's screenshots are placeholders, not images: noted per row, the results file is still written
    assert all("could not be read" in str(r[6]) for r in rows) and not shots._images
    results = [r[2] for r in rows]
    first_pass = results.index("passed") if "passed" in results else len(results)
    assert all(r == "passed" for r in results[first_pass:])          # tests that did not pass come first
    assert str(wb["Waivers"]["A2"].value).startswith("No waivers")   # never a bare header
    assert "Console" not in wb.sheetnames and wb["Screens"]["A2"].value


def test_a_test_result_changed_after_the_seal_is_refused(kit, finished_run):
    run = finished_run
    seal(kit, run)
    progress = kit.json(f"{run}/progress.json")
    progress["tests"]["TC-01"]["result"] = "failed"
    kit.write_json(f"{run}/progress.json", progress)
    r = kit.run(RESULTS, run, expect=1)
    assert "Refused: the sign-off was sealed" in r.out and "progress.json" in r.out


def test_runs_sealed_before_the_fix_are_checked_against_the_sealed_evidence_check(kit, finished_run):
    """Older seals hashed all of progress.json: stage notes alone may differ, test results may not."""
    run = finished_run
    seal(kit, run)
    manifest = kit.json(f"{run}/manifest.json")
    import hashlib
    manifest["files"]["progress.json"] = hashlib.sha256(kit.path(f"{run}/progress.json").read_bytes()).hexdigest()
    kit.write_json(f"{run}/manifest.json", manifest)
    kit.run("scripts/progress.py", "stage", run, "signoff", "started", expect=0)
    kit.run(RESULTS, run, expect=0)
    progress = kit.json(f"{run}/progress.json")
    progress["tests"]["TC-01"]["result"] = "failed"
    kit.write_json(f"{run}/progress.json", progress)
    assert "progress.json" in kit.run(RESULTS, run, expect=1).out
